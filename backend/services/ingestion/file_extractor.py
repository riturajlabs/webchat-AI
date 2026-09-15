"""Text extraction engine for uploaded knowledge files (.txt, .md, .pdf, .docx).

Runs CPU-intensive extraction synchronously so callers can wrap it in
`asyncio.to_thread` without blocking the main event loop. Enforces strict
limits on file size, magic bytes, PDF pages, and extracted characters.
"""

import io
import logging
import zipfile
from dataclasses import dataclass
from pathlib import Path

import docx
import pypdf
import pypdf.errors

from backend.core.errors import (
    AppError,
    DocumentCorruptedError,
    DocumentNoTextError,
    DocumentPasswordProtectedError,
    DocumentTooLargeError,
)

logger = logging.getLogger("webchat_ai")

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_PDF_PAGES = 100
MAX_EXTRACTED_CHARS = 500_000
MIN_EXTRACTED_CHARS = 50

ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

PDF_MAGIC = b"%PDF-"
DOCX_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class ExtractedDocument:
    """Result of file text extraction."""

    filename: str
    text: str
    char_count: int
    mime_type: str
    pages: int | None = None


def _extract_pdf(content: bytes) -> tuple[str, int]:
    """Extract text from a PDF binary payload."""
    if not content.startswith(PDF_MAGIC):
        raise DocumentCorruptedError("Invalid PDF file: missing %PDF- header magic bytes.")

    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocumentPasswordProtectedError("PDF is password-protected or encrypted.")
        page_count = len(reader.pages)
        if page_count > MAX_PDF_PAGES:
            raise DocumentTooLargeError(
                f"PDF exceeds maximum page limit ({page_count} > {MAX_PDF_PAGES} pages)."
            )

        text_parts: list[str] = []
        for page in reader.pages:
            try:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text_parts.append(page_text.strip())
            except pypdf.errors.FileNotDecryptedError as exc:
                raise DocumentPasswordProtectedError(
                    "PDF page is password-protected or encrypted."
                ) from exc
            except Exception as exc:
                logger.warning("pdf_page_extract_warning: %s", exc)

        return "\n\n".join(text_parts), page_count
    except (DocumentPasswordProtectedError, DocumentTooLargeError, DocumentCorruptedError):
        raise
    except pypdf.errors.FileNotDecryptedError as exc:
        raise DocumentPasswordProtectedError("PDF is password-protected or encrypted.") from exc
    except Exception as exc:
        logger.warning("pdf_extraction_error: %s", exc)
        raise DocumentCorruptedError(f"Corrupted or invalid PDF file: {exc}") from exc


def _extract_docx(content: bytes) -> str:
    """Extract text from a Word (.docx) binary payload."""
    if not content.startswith(DOCX_MAGIC):
        raise DocumentCorruptedError("Invalid DOCX file: missing PK zip header magic bytes.")

    try:
        doc = docx.Document(io.BytesIO(content))
        text_parts: list[str] = []

        for p in doc.paragraphs:
            stripped = p.text.strip()
            if stripped:
                text_parts.append(stripped)

        for table in doc.tables:
            for row in table.rows:
                cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells_text:
                    text_parts.append(" | ".join(cells_text))

        return "\n\n".join(text_parts)
    except (zipfile.BadZipFile, Exception) as exc:
        logger.warning("docx_extraction_error: %s", exc)
        raise DocumentCorruptedError(f"Corrupted or invalid DOCX file: {exc}") from exc


def _extract_plain_text(content: bytes) -> str:
    """Extract text from plain text or markdown UTF-8 bytes with fallback."""
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return content.decode("latin-1")
        except Exception as exc:
            raise DocumentCorruptedError(f"Unable to decode text file: {exc}") from exc


def extract_text_from_file(
    filename: str,
    content: bytes,
    mime_type: str | None = None,
) -> ExtractedDocument:
    """Extract readable text from an uploaded file attachment.

    Enforces limits on file size, pages, and extracted text bounds.
    Raises domain AppErrors on any violation.
    """
    file_size = len(content)
    if file_size > MAX_FILE_SIZE_BYTES:
        raise DocumentTooLargeError(
            f"File size {file_size} bytes exceeds maximum allowed limit of "
            f"{MAX_FILE_SIZE_BYTES} bytes (10 MB)."
        )

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        allowed_str = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise AppError(
            f"Unsupported file type '{ext}'. Allowed extensions: {allowed_str}",
            extra={"allowed_extensions": sorted(ALLOWED_EXTENSIONS)},
        )

    pages: int | None = None
    if ext == ".pdf":
        raw_text, pages = _extract_pdf(content)
        resolved_mime = mime_type or "application/pdf"
    elif ext == ".docx":
        raw_text = _extract_docx(content)
        resolved_mime = (
            mime_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    elif ext == ".md":
        raw_text = _extract_plain_text(content)
        resolved_mime = mime_type or "text/markdown"
    else:  # .txt
        raw_text = _extract_plain_text(content)
        resolved_mime = mime_type or "text/plain"

    cleaned_text = raw_text.strip()
    char_count = len(cleaned_text)

    if char_count < MIN_EXTRACTED_CHARS:
        raise DocumentNoTextError(
            f"File '{filename}' contains insufficient readable text "
            f"({char_count} characters; minimum {MIN_EXTRACTED_CHARS} required)."
        )

    if char_count > MAX_EXTRACTED_CHARS:
        raise DocumentTooLargeError(
            f"Extracted text from '{filename}' exceeds maximum allowed limit of "
            f"{MAX_EXTRACTED_CHARS} characters ({char_count} characters)."
        )

    return ExtractedDocument(
        filename=filename,
        text=cleaned_text,
        char_count=char_count,
        mime_type=resolved_mime,
        pages=pages,
    )
